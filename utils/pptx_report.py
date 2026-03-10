"""PPTX report generator with genOway corporate branding."""

import logging
import os
import re
from datetime import datetime

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

import config
from models import Article, EditorReport
from i18n import t

logger = logging.getLogger(__name__)

# ─── genOway brand colors ─────────────────────────────────────────────────────
GENOWAY_RED = RGBColor(0xE3, 0x06, 0x13)       # #E30613
GENOWAY_DARK_RED = RGBColor(0xB0, 0x04, 0x0F)  # #B0040F
GENOWAY_DARK_GRAY = RGBColor(0x33, 0x33, 0x33)  # #333333
GENOWAY_LIGHT_GRAY = RGBColor(0xF2, 0xF2, 0xF2) # #F2F2F2
GENOWAY_WHITE = RGBColor(0xFF, 0xFF, 0xFF)       # #FFFFFF
GENOWAY_MID_GRAY = RGBColor(0x66, 0x66, 0x66)   # #666666

GENOWAY_FONT = "Arial"  # genOway corporate font

SLIDE_WIDTH = Inches(13.333)  # 16:9
SLIDE_HEIGHT = Inches(7.5)


def _add_red_bar(slide, prs):
    """Add the genOway red accent bar at the top of a slide."""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        left=Inches(0), top=Inches(0),
        width=prs.slide_width, height=Inches(0.08),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = GENOWAY_RED
    shape.line.fill.background()


def _add_footer(slide, prs, text="genOway | Confidential"):
    """Add footer text at the bottom of a slide."""
    # Red bottom bar (symmetry with top bar)
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        left=Inches(0), top=prs.slide_height - Inches(0.08),
        width=prs.slide_width, height=Inches(0.08),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = GENOWAY_RED
    shape.line.fill.background()

    # Footer text
    left = Inches(0.5)
    top = prs.slide_height - Inches(0.5)
    width = prs.slide_width - Inches(1)
    height = Inches(0.35)
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(8)
    p.font.name = GENOWAY_FONT
    p.font.color.rgb = GENOWAY_MID_GRAY
    p.alignment = PP_ALIGN.RIGHT


def _set_slide_bg(slide, color=GENOWAY_WHITE):
    """Set solid background color on a slide."""
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def _add_title_slide(prs, title: str, subtitle: str, lang: str):
    """Create the title slide with genOway branding."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
    _set_slide_bg(slide)

    # Red accent block at top
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        left=Inches(0), top=Inches(0),
        width=prs.slide_width, height=Inches(2.5),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = GENOWAY_RED
    shape.line.fill.background()

    # Title text
    txBox = slide.shapes.add_textbox(Inches(1), Inches(0.6), Inches(11), Inches(1.5))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.name = GENOWAY_FONT
    p.font.color.rgb = GENOWAY_WHITE
    p.alignment = PP_ALIGN.LEFT

    # Subtitle
    txBox2 = slide.shapes.add_textbox(Inches(1), Inches(3.0), Inches(11), Inches(1))
    tf2 = txBox2.text_frame
    tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    p2.text = subtitle
    p2.font.size = Pt(18)
    p2.font.name = GENOWAY_FONT
    p2.font.color.rgb = GENOWAY_DARK_GRAY
    p2.alignment = PP_ALIGN.LEFT

    # Date
    txBox3 = slide.shapes.add_textbox(Inches(1), Inches(4.2), Inches(11), Inches(0.5))
    tf3 = txBox3.text_frame
    p3 = tf3.paragraphs[0]
    p3.text = datetime.now().strftime("%d/%m/%Y")
    p3.font.size = Pt(14)
    p3.font.name = GENOWAY_FONT
    p3.font.color.rgb = GENOWAY_MID_GRAY

    # genOway branding
    txBox4 = slide.shapes.add_textbox(Inches(1), Inches(6.5), Inches(4), Inches(0.5))
    tf4 = txBox4.text_frame
    p4 = tf4.paragraphs[0]
    run = p4.add_run()
    run.text = "genOway"
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.name = GENOWAY_FONT
    run.font.color.rgb = GENOWAY_RED


def _add_section_slide(prs, section_title: str):
    """Add a section divider slide."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_bg(slide, GENOWAY_LIGHT_GRAY)

    # Red left bar
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        left=Inches(0), top=Inches(0),
        width=Inches(0.15), height=prs.slide_height,
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = GENOWAY_RED
    shape.line.fill.background()

    # Section title
    txBox = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(10), Inches(2))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = section_title
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.name = GENOWAY_FONT
    p.font.color.rgb = GENOWAY_DARK_GRAY
    p.alignment = PP_ALIGN.LEFT

    _add_footer(slide, prs)


def _add_content_slide(prs, title: str, body_text: str):
    """Add a standard content slide with title and body text."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_bg(slide)
    _add_red_bar(slide, prs)

    # Title
    txBox = slide.shapes.add_textbox(Inches(0.7), Inches(0.3), Inches(11.5), Inches(0.9))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.name = GENOWAY_FONT
    p.font.color.rgb = GENOWAY_RED
    p.alignment = PP_ALIGN.LEFT

    # Body
    txBox2 = slide.shapes.add_textbox(Inches(0.7), Inches(1.4), Inches(11.5), Inches(5.5))
    tf2 = txBox2.text_frame
    tf2.word_wrap = True

    # Split body into lines and format
    lines = body_text.strip().split("\n")
    for idx, line in enumerate(lines):
        line = line.strip()
        if not line:
            continue
        if idx == 0:
            p = tf2.paragraphs[0]
        else:
            p = tf2.add_paragraph()

        # Bullet point detection
        if line.startswith("- ") or line.startswith("* "):
            line = line[2:]
            p.level = 0
            p.space_before = Pt(4)

        # Bold detection for **text**
        parts = re.split(r'\*\*(.+?)\*\*', line)
        for i_part, part in enumerate(parts):
            if not part:
                continue
            run = p.add_run()
            run.text = part
            run.font.size = Pt(14)
            run.font.name = GENOWAY_FONT
            run.font.color.rgb = GENOWAY_DARK_GRAY
            if i_part % 2 == 1:  # odd parts are bold
                run.font.bold = True

    _add_footer(slide, prs)


def _add_figure_slide(prs, figure_path: str, caption: str):
    """Add a slide with a figure image."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_bg(slide)
    _add_red_bar(slide, prs)

    if os.path.exists(figure_path):
        # Center the image
        img_width = Inches(10)
        img_left = (prs.slide_width - img_width) // 2
        slide.shapes.add_picture(
            figure_path, img_left, Inches(0.5),
            width=img_width,
        )

    # Caption
    txBox = slide.shapes.add_textbox(Inches(1), Inches(6.5), Inches(11), Inches(0.5))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = caption
    p.font.size = Pt(11)
    p.font.italic = True
    p.font.name = GENOWAY_FONT
    p.font.color.rgb = GENOWAY_MID_GRAY
    p.alignment = PP_ALIGN.CENTER

    _add_footer(slide, prs)


def _add_kpi_slide(prs, article_count: int, sources: str, confidence: float,
                   year_range: str, lang: str):
    """Add a KPI overview slide."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_bg(slide)
    _add_red_bar(slide, prs)

    # Title
    txBox = slide.shapes.add_textbox(Inches(0.7), Inches(0.3), Inches(11.5), Inches(0.9))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = "Key Metrics" if lang == "en" else "Indicateurs cles"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.name = GENOWAY_FONT
    p.font.color.rgb = GENOWAY_RED

    # KPI cards
    kpis = [
        (str(article_count), t("html_articles_analyzed", lang)),
        (sources, "Sources"),
    ]
    if confidence > 0:
        kpis.append((f"{confidence:.0f}%", t("html_factual_confidence", lang)))
    if year_range:
        kpis.append((year_range, t("html_period_covered", lang)))

    card_width = Inches(2.8)
    card_height = Inches(2.5)
    gap = Inches(0.3)
    total_width = len(kpis) * card_width + (len(kpis) - 1) * gap
    start_left = (prs.slide_width - total_width) // 2

    for i, (value, label) in enumerate(kpis):
        left = start_left + i * (card_width + gap)
        top = Inches(2.2)

        # Card background
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            left=left, top=top, width=card_width, height=card_height,
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = GENOWAY_LIGHT_GRAY
        shape.line.fill.background()

        # Value
        txVal = slide.shapes.add_textbox(left, top + Inches(0.4), card_width, Inches(1.2))
        tf_val = txVal.text_frame
        tf_val.word_wrap = True
        p_val = tf_val.paragraphs[0]
        p_val.alignment = PP_ALIGN.CENTER
        run_val = p_val.add_run()
        run_val.text = value
        run_val.font.size = Pt(32)
        run_val.font.bold = True
        run_val.font.name = GENOWAY_FONT
        run_val.font.color.rgb = GENOWAY_RED

        # Label
        txLbl = slide.shapes.add_textbox(left, top + Inches(1.6), card_width, Inches(0.6))
        tf_lbl = txLbl.text_frame
        tf_lbl.word_wrap = True
        p_lbl = tf_lbl.paragraphs[0]
        p_lbl.alignment = PP_ALIGN.CENTER
        run_lbl = p_lbl.add_run()
        run_lbl.text = label.upper()
        run_lbl.font.size = Pt(10)
        run_lbl.font.bold = True
        run_lbl.font.name = GENOWAY_FONT
        run_lbl.font.color.rgb = GENOWAY_MID_GRAY

    _add_footer(slide, prs)


def _split_report_into_sections(report_markdown: str) -> list[tuple[str, str]]:
    """Split markdown report into (title, body) sections based on ## headers."""
    sections = []
    current_title = ""
    current_body_lines = []

    for line in report_markdown.split("\n"):
        if line.startswith("## "):
            if current_title or current_body_lines:
                sections.append((current_title, "\n".join(current_body_lines)))
            current_title = line.lstrip("#").strip()
            current_body_lines = []
        elif line.startswith("# ") and not line.startswith("##"):
            if current_title or current_body_lines:
                sections.append((current_title, "\n".join(current_body_lines)))
            current_title = line.lstrip("#").strip()
            current_body_lines = []
        else:
            current_body_lines.append(line)

    if current_title or current_body_lines:
        sections.append((current_title, "\n".join(current_body_lines)))

    return sections


def _paginate_text(text: str, max_chars: int = 1200) -> list[str]:
    """Split text into pages that fit on a slide."""
    lines = text.strip().split("\n")
    pages = []
    current_page = []
    current_len = 0

    for line in lines:
        line_len = len(line) + 1
        if current_len + line_len > max_chars and current_page:
            pages.append("\n".join(current_page))
            current_page = []
            current_len = 0
        current_page.append(line)
        current_len += line_len

    if current_page:
        pages.append("\n".join(current_page))

    return pages


def generate_pptx(
    report_markdown: str,
    articles: list[Article],
    editor_report: EditorReport,
    keywords: str,
    date_from: str,
    date_to: str,
    lang: str = "fr",
    num_slides: int = 15,
    template_path: str | None = None,
) -> str:
    """Generate a PPTX presentation with genOway branding.

    Args:
        report_markdown: The markdown report
        articles: List of articles
        editor_report: Editor report with figures and verifications
        keywords: Search keywords
        date_from: Start date
        date_to: End date
        lang: Language
        num_slides: Target number of slides (approximate)
        template_path: Optional path to a blank PPTX template to use as base

    Returns:
        Path to the generated PPTX file
    """
    if template_path and os.path.exists(template_path):
        prs = Presentation(template_path)
        # Remove existing slides from the template (keep masters/theme only)
        while len(prs.slides) > 0:
            rId = prs.slides._sldIdLst[0].get('r:id')
            prs.part.drop_rel(rId)
            prs.slides._sldIdLst.remove(prs.slides._sldIdLst[0])
    else:
        prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT

    # 1. Title slide
    report_title = t("html_title", lang)
    subtitle = f"{keywords}\n{date_from} — {date_to}"
    _add_title_slide(prs, report_title, subtitle, lang)

    # 2. KPI slide
    sources_list = set(a.source for a in articles)
    sources_str = ", ".join(sorted(sources_list))
    years = [a.year for a in articles if a.year]
    year_range = f"{min(years)}-{max(years)}" if years else ""
    _add_kpi_slide(prs, len(articles), sources_str,
                   editor_report.confidence_score, year_range, lang)

    # 3. Report content slides
    sections = _split_report_into_sections(report_markdown)

    # Budget remaining slides: title(1) + kpi(1) + figures + closing(1)
    figure_count = len([f for f in editor_report.figures_generated if os.path.exists(f)])
    slides_for_content = max(num_slides - 2 - figure_count - 1, len(sections))

    # Determine max chars per slide based on available space
    total_text_len = sum(len(body) for _, body in sections)
    if slides_for_content > 0 and total_text_len > 0:
        max_chars_per_slide = max(800, total_text_len // slides_for_content)
    else:
        max_chars_per_slide = 1200

    for section_title, section_body in sections:
        if not section_body.strip():
            continue

        pages = _paginate_text(section_body, max_chars=max_chars_per_slide)

        for page_idx, page_text in enumerate(pages):
            slide_title = section_title
            if len(pages) > 1:
                slide_title = f"{section_title} ({page_idx + 1}/{len(pages)})"
            _add_content_slide(prs, slide_title, page_text)

    # 4. Figure slides
    figure_caption_map = {
        "timeline.png": t("fig_timeline", lang),
        "top_journals.png": t("fig_top_journals", lang),
        "wordcloud.png": t("fig_wordcloud", lang),
        "cocitation_network.png": t("fig_cocitation", lang),
        "theme_heatmap.png": t("fig_heatmap", lang),
    }

    if editor_report.figures_generated:
        _add_section_slide(prs, t("html_infographics", lang))
        for fig_path in editor_report.figures_generated:
            if os.path.exists(fig_path):
                fname = os.path.basename(fig_path)
                caption = figure_caption_map.get(fname, fname)
                _add_figure_slide(prs, fig_path, caption)

    # 5. Closing slide
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_bg(slide, GENOWAY_RED)

    txBox = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11), Inches(2))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = "Thank you" if lang == "en" else "Merci"
    run.font.size = Pt(44)
    run.font.bold = True
    run.font.name = GENOWAY_FONT
    run.font.color.rgb = GENOWAY_WHITE

    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    run2 = p2.add_run()
    run2.text = "genOway"
    run2.font.size = Pt(28)
    run2.font.bold = True
    run2.font.name = GENOWAY_FONT
    run2.font.color.rgb = GENOWAY_WHITE

    p3 = tf.add_paragraph()
    p3.alignment = PP_ALIGN.CENTER
    p3.space_before = Pt(20)
    run3 = p3.add_run()
    run3.text = "www.genoway.com"
    run3.font.size = Pt(14)
    run3.font.name = GENOWAY_FONT
    run3.font.color.rgb = GENOWAY_WHITE

    # Save
    os.makedirs(config.OUTPUT_DIR, exist_ok=True)
    output_path = os.path.join(config.OUTPUT_DIR, "rapport_bibliographique.pptx")
    prs.save(output_path)

    logger.info(f"PPTX report generated: {output_path} ({len(prs.slides)} slides)")
    return output_path
